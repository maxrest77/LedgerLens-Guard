from backend.utils import utc_now
import io
import json
import pytest
from datetime import datetime
from fastapi.testclient import TestClient
from sqlmodel import Session, SQLModel, create_engine, select
from sqlalchemy.pool import StaticPool

from backend.api.main import app
from backend.api.auth import get_db, get_password_hash
from backend.data.schema import (
    ReconciliationCase, CaseStatus, Reviewer, Role, ApprovalRequest, EvidenceAttachment
)
from backend.audit.chain import AuditBlock
from backend.engine.sanitizer import is_luhn_valid, mask_pan_luhn
from backend.engine.evidence_parser import (
    route_and_parse, parse_delimited, parse_camt053, parse_mt940, parse_bai2, parse_xlsx, parse_document_text
)

# ── Test Sanitizer (PCI-DSS & Luhn Checksum) ──────────────────────────────────

def test_luhn_validation_and_masking():
    # Valid Visa card number test pattern (passes Luhn)
    valid_visa = "4532015012345671"
    assert is_luhn_valid(valid_visa) is True

    # Invalid card number (fails Luhn)
    invalid_num = "4532015012345672"
    assert is_luhn_valid(invalid_num) is False

    # Text containing valid PAN should be masked
    sample_text = f"Payment captured using card {valid_visa} on gateway."
    masked = mask_pan_luhn(sample_text)
    assert valid_visa not in masked
    assert "XXXX-XXXX-XXXX-5671" in masked

    # Text containing invalid number should remain unmasked
    sample_invalid = f"Order reference {invalid_num} updated."
    assert mask_pan_luhn(sample_invalid) == sample_invalid

# ── Test Multi-Format Parsers ──────────────────────────────────────────────────

def test_csv_parser_fuzzy_headers_and_integer_paisa():
    csv_data = """UTR,Amount,Fee,Tax,Net_Amount,Date,Customer_Card
UTR_TEST_1001,1000.50,20.00,3.60,976.90,2026-09-01,4532-0150-1234-5671
UTR_TEST_1002,500.00,10.00,1.80,488.20,2026-09-01,9876543210987654
"""
    result = parse_delimited(csv_data.encode("utf-8"), "settlement.csv")
    recs = result["records"]
    assert len(recs) == 2
    
    # Check integer paisa conversions
    assert recs[0]["utr"] == "UTR_TEST_1001"
    assert recs[0]["amount_paisa"] == 100050  # 1000.50 * 100
    assert recs[0]["fee_paisa"] == 2000       # 20.00 * 100
    assert recs[0]["tax_paisa"] == 360        # 3.60 * 100
    assert recs[0]["net_paisa"] == 97690      # 976.90 * 100
    
    # Check PCI-DSS card masking
    assert "XXXX-XXXX-XXXX-5671" in recs[0]["masked_account_or_pan"]
    assert result["summary"]["pci_masked_count"] >= 1

def test_iso20022_camt053_parser():
    camt_xml = """<?xml version="1.0" encoding="UTF-8"?>
<Document xmlns="urn:iso:std:iso:20022:tech:xsd:camt.053.001.02">
  <BkToCstmrStmt>
    <Stmt>
      <Ntry>
        <Amt Ccy="INR">15000.00</Amt>
        <CdtDbtInd>CRDT</CdtDbtInd>
        <BookgDt><Dt>2026-09-02</Dt></BookgDt>
        <AcctSvcrRef>AURA_CAMT_9988</AcctSvcrRef>
        <AddtlNtryInf>Settlement credit for merchant M123</AddtlNtryInf>
      </Ntry>
    </Stmt>
  </BkToCstmrStmt>
</Document>"""
    result = parse_camt053(camt_xml.encode("utf-8"))
    assert len(result["records"]) == 1
    rec = result["records"][0]
    assert rec["utr"] == "AURA_CAMT_9988"
    assert rec["amount_paisa"] == 1500000
    assert rec["net_paisa"] == 1500000

def test_camt053_endtoend_id_and_namespaces():
    camt_xml = """<?xml version="1.0" encoding="UTF-8"?>
<camt:Document xmlns:camt="urn:iso:std:iso:20022:tech:xsd:camt.053.001.08">
  <camt:BkToCstmrStmt>
    <camt:Stmt>
      <camt:Ntry>
        <camt:Amt Ccy="USD">4500.00</camt:Amt>
        <camt:CdtDbtInd>CRDT</camt:CdtDbtInd>
        <camt:BookgDt><camt:DtTm>2026-09-02T14:20:00</camt:DtTm></camt:BookgDt>
        <camt:NtryDtls>
          <camt:TxDtls>
            <camt:Refs>
              <camt:EndToEndId>E2E_PAYPAL_991122</camt:EndToEndId>
              <camt:TxId>TX_BANK_554433</camt:TxId>
            </camt:Refs>
            <camt:Amt Ccy="USD">4500.00</camt:Amt>
            <camt:RmtInf><camt:Ustrd>Online payment settlement</camt:Ustrd></camt:RmtInf>
          </camt:TxDtls>
        </camt:NtryDtls>
      </camt:Ntry>
    </camt:Stmt>
  </camt:BkToCstmrStmt>
</camt:Document>"""
    result = route_and_parse(camt_xml.encode("utf-8"), "statement.xml")
    assert result["file_type"] == "CAMT053"
    assert len(result["records"]) == 1
    rec = result["records"][0]
    assert rec["utr"] == "E2E_PAYPAL_991122"
    assert rec["amount_paisa"] == 450000
    assert rec["net_paisa"] == 450000

def test_camt053_bypasses_notprovided_to_clrsysref_with_fees():
    camt_xml = """<?xml version="1.0" encoding="UTF-8"?>
<Document xmlns="urn:iso:std:iso:20022:tech:xsd:camt.053.001.04">
  <BkToCstmrStmt>
    <Stmt>
      <Ntry>
        <Amt Ccy="EUR">800.00</Amt>
        <CdtDbtInd>DBIT</CdtDbtInd>
        <BookgDt><Dt>2026-09-01</Dt></BookgDt>
        <NtryDtls>
          <TxDtls>
            <Refs>
              <EndToEndId>NOTPROVIDED</EndToEndId>
              <ClrSysRef>SEPA_CLEAR_44556677</ClrSysRef>
            </Refs>
            <Amt Ccy="EUR">800.00</Amt>
            <Chrgs><Amt Ccy="EUR">12.00</Amt></Chrgs>
          </TxDtls>
        </NtryDtls>
      </Ntry>
    </Stmt>
  </BkToCstmrStmt>
</Document>"""
    result = parse_camt053(camt_xml.encode("utf-8"))
    assert len(result["records"]) == 1
    rec = result["records"][0]
    assert rec["utr"] == "SEPA_CLEAR_44556677"
    assert rec["amount_paisa"] == 80000
    assert rec["fee_paisa"] == 1200
    assert rec["net_paisa"] == 78800

def test_camt053_batch_booking_multiple_txdtls():
    camt_xml = """<?xml version="1.0" encoding="UTF-8"?>
<Document xmlns="urn:iso:std:iso:20022:tech:xsd:camt.053.001.02">
  <BkToCstmrStmt>
    <Stmt>
      <Ntry>
        <Amt Ccy="INR">3000.00</Amt>
        <CdtDbtInd>CRDT</CdtDbtInd>
        <BookgDt><Dt>2026-09-02</Dt></BookgDt>
        <NtryDtls>
          <TxDtls>
            <Refs><InstrId>INSTR_PART_1</InstrId></Refs>
            <Amt Ccy="INR">1000.00</Amt>
            <Chrgs><Amt Ccy="INR">20.00</Amt></Chrgs>
          </TxDtls>
          <TxDtls>
            <Refs><InstrId>INSTR_PART_2</InstrId></Refs>
            <Amt Ccy="INR">2000.00</Amt>
            <Chrgs><Amt Ccy="INR">40.00</Amt></Chrgs>
          </TxDtls>
        </NtryDtls>
      </Ntry>
    </Stmt>
  </BkToCstmrStmt>
</Document>"""
    result = parse_camt053(camt_xml.encode("utf-8"))
    assert len(result["records"]) == 2
    rec1 = result["records"][0]
    assert rec1["utr"] == "INSTR_PART_1"
    assert rec1["amount_paisa"] == 100000
    assert rec1["fee_paisa"] == 2000
    assert rec1["net_paisa"] == 98000
    rec2 = result["records"][1]
    assert rec2["utr"] == "INSTR_PART_2"
    assert rec2["amount_paisa"] == 200000
    assert rec2["fee_paisa"] == 4000
    assert rec2["net_paisa"] == 196000
    assert result["summary"]["total_amount_paisa"] == 300000
    assert result["summary"]["net_amount_paisa"] == 294000

def test_camt053_structured_remittance_creditor_ref():
    camt_xml = """<?xml version="1.0" encoding="UTF-8"?>
<Document>
  <BkToCstmrStmt>
    <Stmt>
      <Ntry>
        <Amt Ccy="INR">1200.00</Amt>
        <BookgDt><Dt>2026-09-02</Dt></BookgDt>
        <NtryDtls>
          <TxDtls>
            <RmtInf>
              <Strd>
                <CdtrRefInf>
                  <Ref>INV_CRED_998877</Ref>
                </CdtrRefInf>
              </Strd>
            </RmtInf>
            <Amt Ccy="INR">1200.00</Amt>
          </TxDtls>
        </NtryDtls>
      </Ntry>
    </Stmt>
  </BkToCstmrStmt>
</Document>"""
    result = parse_camt053(camt_xml.encode("utf-8"))
    assert len(result["records"]) == 1
    assert result["records"][0]["utr"] == "INV_CRED_998877"
    assert result["records"][0]["amount_paisa"] == 120000

def test_camt053_narrative_reference_extraction():
    camt_xml = """<?xml version="1.0" encoding="UTF-8"?>
<Document>
  <BkToCstmrStmt>
    <Stmt>
      <Ntry>
        <Amt Ccy="INR">500.00</Amt>
        <BookgDt><Dt>2026-09-02</Dt></BookgDt>
        <AddtlNtryInf>NEFT Inward UTR: PUNBH20260902123456 - Merchant settlement</AddtlNtryInf>
      </Ntry>
    </Stmt>
  </BkToCstmrStmt>
</Document>"""
    result = parse_camt053(camt_xml.encode("utf-8"))
    assert len(result["records"]) == 1
    assert result["records"][0]["utr"] == "PUNBH20260902123456"
    assert result["records"][0]["amount_paisa"] == 50000

def test_camt053_last_resort_fallback_camt_idx():
    camt_xml = """<?xml version="1.0" encoding="UTF-8"?>
<Document>
  <BkToCstmrStmt>
    <Stmt>
      <Ntry>
        <Amt Ccy="INR">100.00</Amt>
        <BookgDt><Dt>2026-09-02</Dt></BookgDt>
      </Ntry>
    </Stmt>
  </BkToCstmrStmt>
</Document>"""
    result = parse_camt053(camt_xml.encode("utf-8"))
    assert len(result["records"]) == 1
    assert result["records"][0]["utr"] == "CAMT_1"
    assert result["records"][0]["amount_paisa"] == 10000

def test_camt053_upload_and_executive_pack_e2e(client, engine):
    with Session(engine) as session:
        rev = Reviewer(email="reviewer_camt@test.com", hashed_password=get_password_hash("pw"), role=Role.REVIEWER, portfolio_id="PORTFOLIO_A")
        case = ReconciliationCase(
            case_id="case_camt_e2e",
            portfolio_id="PORTFOLIO_A",
            exception_code="BANK_CREDIT_SHORTFALL",
            severity="HIGH",
            status=CaseStatus.OPEN,
            expected_paisa=5000000,
            actual_paisa=4500000,
            delta_paisa=500000,
            confidence_score=0.96,
            opened_at=utc_now(),
            explanation="Bank settlement credit difference",
            suggested_action="Review ISO 20022 statement"
        )
        session.add_all([rev, case])
        session.commit()

    # Login as reviewer
    res_login = client.post("/auth/login", data={"username": "reviewer_camt@test.com", "password": "pw"})
    cookie = res_login.cookies.get("session_token")

    camt_xml_bytes = b"""<?xml version="1.0" encoding="UTF-8"?>
<Document xmlns="urn:iso:std:iso:20022:tech:xsd:camt.053.001.08">
  <BkToCstmrStmt>
    <Stmt>
      <Ntry>
        <Amt Ccy="INR">45000.00</Amt>
        <CdtDbtInd>CRDT</CdtDbtInd>
        <BookgDt><Dt>2026-09-02</Dt></BookgDt>
        <NtryDtls>
          <TxDtls>
            <Refs><EndToEndId>E2E_CAMT_RESOLVED_9999</EndToEndId></Refs>
            <Amt Ccy="INR">45000.00</Amt>
            <RmtInf><Ustrd>Settlement credit from gateway</Ustrd></RmtInf>
          </TxDtls>
        </NtryDtls>
      </Ntry>
    </Stmt>
  </BkToCstmrStmt>
</Document>"""

    # Upload CAMT.053 XML evidence
    res_upload = client.post(
        "/api/exceptions/case_camt_e2e/evidence/upload",
        files={"file": ("bank_statement_iso20022.xml", io.BytesIO(camt_xml_bytes), "application/xml")},
        headers={"X-CSRF-Protection": "1"},
        cookies={"session_token": cookie}
    )
    assert res_upload.status_code == 200
    upload_data = res_upload.json()
    assert upload_data["status"] == "staged"
    assert upload_data["file_type"] == "CAMT053"
    assert upload_data["summary"]["total_records"] == 1
    assert upload_data["summary"]["total_amount_paisa"] == 4500000

    recs = upload_data["extracted_records"]
    assert len(recs) == 1
    assert recs[0]["utr"] == "E2E_CAMT_RESOLVED_9999"
    assert recs[0]["amount_paisa"] == 4500000

    # Submit review to permanently commit proof
    res_review = client.post(
        "/api/exceptions/case_camt_e2e/review",
        json={"action": "APPROVE", "reason": "Verified ISO 20022 CAMT.053 bank statement with EndToEndId reference."},
        headers={"X-CSRF-Protection": "1", "Idempotency-Key": "test-idem-camt-1"},
        cookies={"session_token": cookie}
    )
    assert res_review.status_code == 200

    # Verify Executive Pack has committed CAMT053 file and extracted records
    res_exec = client.get("/api/exceptions/case_camt_e2e/executive-pack", cookies={"session_token": cookie})
    assert res_exec.status_code == 200
    pack = res_exec.json()
    attached_files = pack["evidence_pack"]["attached_evidence_files"]
    assert len(attached_files) == 1
    assert attached_files[0]["file_type"] == "CAMT053"
    assert attached_files[0]["filename"] == "bank_statement_iso20022.xml"
    extracted_txns = pack["evidence_pack"]["extracted_transactions"]
    assert len(extracted_txns) == 1
    assert extracted_txns[0]["utr"] == "E2E_CAMT_RESOLVED_9999"
    assert extracted_txns[0]["amount_paisa"] == 4500000

def test_swift_mt940_parser():
    mt940_txt = """:20:STMT20260901
:25:1234567890
:28C:001
:60F:C260901INR100000,00
:61:2609010901CR5000,00NTRFNONREF//NOVA_UTR_4455
:86:UPI SETTLEMENT CREDIT UTR: NOVA_UTR_4455
:62F:C260901INR105000,00"""
    result = parse_mt940(mt940_txt)
    assert len(result["records"]) == 1
    rec = result["records"][0]
    assert "NOVA_UTR_4455" in rec["utr"]
    assert rec["amount_paisa"] == 500000

def test_bai2_parser():
    bai2_txt = """01,SENDER,RECEIVER,260901,1200,1,80,1/
02,123456,ABC,1,260901,1200,,2/
03,987654321,INR/
16,195,750000,Z,TRUE_UTR_8899,,UPI MERCHANT PAYOUT/
49,750000,1/
98,750000,1,4/
99,750000,1,6/"""
    result = parse_bai2(bai2_txt)
    assert len(result["records"]) == 1
    rec = result["records"][0]
    assert rec["utr"] == "TRUE_UTR_8899"
    assert rec["amount_paisa"] == 75000000

# ── Test Endpoints & Audit Hash Chaining ───────────────────────────────────────

@pytest.fixture(name="engine")
def engine_fixture():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool
    )
    SQLModel.metadata.create_all(engine)
    return engine

@pytest.fixture(name="client")
def client_fixture(engine):
    app.dependency_overrides[get_db] = lambda: Session(engine)
    client = TestClient(app)
    yield client
    app.dependency_overrides.clear()

def test_evidence_upload_and_executive_pack_e2e(client, engine):
    with Session(engine) as session:
        # Create test users
        rev = Reviewer(email="reviewer@test.com", hashed_password=get_password_hash("pw"), role=Role.REVIEWER, portfolio_id="PORTFOLIO_A")
        senior = Reviewer(email="admin@test.com", hashed_password=get_password_hash("pw"), role=Role.ADMIN, portfolio_id="PORTFOLIO_A")
        
        # Create case in PORTFOLIO_A
        case = ReconciliationCase(
            case_id="case_ev_1",
            portfolio_id="PORTFOLIO_A",
            exception_code="BANK_CREDIT_SHORTFALL",
            severity="CRITICAL",
            status=CaseStatus.OPEN,
            expected_paisa=100000,
            actual_paisa=70000,
            delta_paisa=30000,
            confidence_score=0.92,
            opened_at=utc_now(),
            explanation="Bank deposit was less than expected settlement net",
            suggested_action="Request bank trace advice slip"
        )
        session.add_all([rev, senior, case])
        session.commit()

    # Login as reviewer
    res_login = client.post("/auth/login", data={"username": "reviewer@test.com", "password": "pw"})
    cookie = res_login.cookies.get("session_token")

    # 1. Upload CSV evidence (Stages evidence, does not pollute audit chain yet)
    csv_file_bytes = b"UTR,Amount,Fee,Tax,Net_Amount,Date\nUTR_BANK_EVID_1,300.00,0,0,300.00,2026-09-02\n"
    res_upload = client.post(
        "/api/exceptions/case_ev_1/evidence/upload",
        files={"file": ("advice_slip.csv", io.BytesIO(csv_file_bytes), "text/csv")},
        headers={"X-CSRF-Protection": "1"},
        cookies={"session_token": cookie}
    )
    assert res_upload.status_code == 200
    upload_body = res_upload.json()
    assert upload_body["status"] == "staged"
    assert upload_body["is_committed"] is False
    assert upload_body["file_type"] == "CSV"
    assert upload_body["summary"]["total_records"] == 1
    assert upload_body["summary"]["total_amount_paisa"] == 30000

    # Before review submission, audit block is NOT yet chained
    with Session(engine) as session:
        blocks = session.exec(select(AuditBlock).where(AuditBlock.case_id == "case_ev_1")).all()
        assert len(blocks) == 0

    # Default evidence list only shows committed proofs (should be 0 right now)
    res_list_committed = client.get("/api/exceptions/case_ev_1/evidence", cookies={"session_token": cookie})
    assert res_list_committed.status_code == 200
    assert len(res_list_committed.json()["data"]) == 0

    # Query with include_staged=True shows the staged draft
    res_list_staged = client.get("/api/exceptions/case_ev_1/evidence?include_staged=true", cookies={"session_token": cookie})
    assert res_list_staged.status_code == 200
    staged_items = res_list_staged.json()["data"]
    assert len(staged_items) == 1
    assert staged_items[0]["filename"] == "advice_slip.csv"
    assert staged_items[0]["is_committed"] is False

    # 2. Submit Maker review with reason (This permanently commits the staged proof and seals the audit block)
    res_review = client.post(
        "/api/exceptions/case_ev_1/review",
        json={"action": "APPROVE", "reason": "Verified bank credit advice slip attached for this case settlement."},
        headers={"X-CSRF-Protection": "1", "Idempotency-Key": "test-idem-evidence-1"},
        cookies={"session_token": cookie}
    )
    assert res_review.status_code == 200

    # Check that audit chain now contains EVIDENCE_ATTACHED block followed by PROPOSE_APPROVE
    with Session(engine) as session:
        blocks = session.exec(select(AuditBlock).where(AuditBlock.case_id == "case_ev_1").order_by(AuditBlock.index.asc())).all()
        assert len(blocks) >= 2
        assert blocks[0].action == "EVIDENCE_ATTACHED"
        assert "advice_slip.csv" in blocks[0].reason

    # Default evidence list now contains the committed attachment
    res_list = client.get("/api/exceptions/case_ev_1/evidence", cookies={"session_token": cookie})
    assert res_list.status_code == 200
    attachments = res_list.json()["data"]
    assert len(attachments) == 1
    assert attachments[0]["filename"] == "advice_slip.csv"
    assert attachments[0]["is_committed"] is True

    # 3. Fetch Executive Pack Payload (Shows permanently committed evidence)
    res_exec = client.get("/api/exceptions/case_ev_1/executive-pack", cookies={"session_token": cookie})
    assert res_exec.status_code == 200
    pack = res_exec.json()
    
    assert pack["case"]["case_id"] == "case_ev_1"
    assert pack["ledger"]["expected_paisa"] == 100000
    assert pack["ledger"]["actual_paisa"] == 70000
    assert pack["ledger"]["delta_paisa"] == 30000
    assert pack["governance"]["maker"]["email"] == "reviewer@test.com"
    assert pack["governance"]["dual_control_enforced"] is True
    assert len(pack["evidence_pack"]["attached_evidence_files"]) == 1
    assert pack["cryptographic_seal"]["chain_valid"] is True
    assert pack["cryptographic_seal"]["latest_block_hash"] is not None


# ── Test Dynamic XLSX Spreadsheet Parser ───────────────────────────────────────

def test_xlsx_parser_standard_headers_and_integer_paisa():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(["Transaction ID", "Amount", "Fee", "Tax", "Net Amount", "Date", "Customer Card"])
    ws.append(["TXN_XLSX_101", 2500.50, 50.00, 9.00, 2441.50, "2026-09-02", "4532015012345671"])
    ws.append(["TXN_XLSX_102", 1000.00, 20.00, 3.60, 976.40, "2026-09-02", "9876543210987654"])

    buf = io.BytesIO()
    wb.save(buf)
    result = route_and_parse(buf.getvalue(), "settlement_batch.xlsx")

    assert result["file_type"] == "XLSX"
    recs = result["records"]
    assert len(recs) == 2
    assert recs[0]["utr"] == "TXN_XLSX_101"
    assert recs[0]["amount_paisa"] == 250050
    assert recs[0]["fee_paisa"] == 5000
    assert recs[0]["tax_paisa"] == 900
    assert recs[0]["net_paisa"] == 244150
    assert "XXXX-XXXX-XXXX-5671" in recs[0]["masked_account_or_pan"]
    assert result["summary"]["total_amount_paisa"] == 350050
    assert result["summary"]["net_amount_paisa"] == 341790
    assert result["summary"]["pci_masked_count"] >= 1


def test_xlsx_parser_offset_headers_and_formatted_currencies():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bank Statement"
    ws.append(["STATE BANK OF INDIA - MERCHANT SETTLEMENT REPORT"])
    ws.append(["Account: 1234567890", "Period: SEP 2026"])
    ws.append([])  # blank row
    ws.append(["UTR / Reference", "Value Date", "Credit Amount (INR)", "MDR Charges", "GST", "Net Payout"])
    ws.append(["SBI_UTR_8899001", datetime(2026, 9, 2), "₹ 50,000.00", "₹ 1,000.00", "₹ 180.00", "₹ 48,820.00"])
    ws.append(["Total", None, 50000.00, 1000.00, 180.00, 48820.00])  # footer summary should be omitted

    buf = io.BytesIO()
    wb.save(buf)
    result = route_and_parse(buf.getvalue(), "sbi_advice.xlsx")

    assert result["file_type"] == "XLSX"
    recs = result["records"]
    assert len(recs) == 1
    assert recs[0]["utr"] == "SBI_UTR_8899001"
    assert recs[0]["amount_paisa"] == 5000000
    assert recs[0]["fee_paisa"] == 100000
    assert recs[0]["tax_paisa"] == 18000
    assert recs[0]["net_paisa"] == 4882000
    assert recs[0]["timestamp"] == "2026-09-02"


def test_xlsx_parser_multiple_sheets():
    import openpyxl
    wb = openpyxl.Workbook()
    ws_cover = wb.active
    ws_cover.title = "Cover Page"
    ws_cover.append(["Document Information"])
    ws_cover.append(["Author", "Finance Ops"])

    ws_data = wb.create_sheet(title="Settlement_Details")
    ws_data.append(["Ref No", "Paid Amount", "Settled Date"])
    ws_data.append(["VELOCEPAY_REF_4411", 750.25, "2026-09-02"])

    buf = io.BytesIO()
    wb.save(buf)
    result = route_and_parse(buf.getvalue(), "multi_sheet_report.xlsx")

    assert result["file_type"] == "XLSX"
    recs = result["records"]
    assert len(recs) == 1
    assert recs[0]["utr"] == "VELOCEPAY_REF_4411"
    assert recs[0]["amount_paisa"] == 75025
    assert recs[0]["net_paisa"] == 75025
    assert recs[0]["timestamp"] == "2026-09-02"


def test_xlsx_parser_heuristics_and_narrative_recovery():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Adjustments"
    ws.append(["Date", "Description / Particulars", "Amount (INR)", "Fee"])
    ws.append(["2026-09-02", "Reversal for dispute UTR_REV_99887766 on UPI gateway", "(1,500.50)", 0])
    ws.append(["2026-09-02", "Credit refund REF_ABC_12345678", "250.00 CR", "5.00"])

    buf = io.BytesIO()
    wb.save(buf)
    result = route_and_parse(buf.getvalue(), "adjustments.xlsx")

    assert result["file_type"] == "XLSX"
    recs = result["records"]
    assert len(recs) == 2
    assert recs[0]["amount_paisa"] == -150050
    assert "UTR_REV_99887766" in recs[0]["utr"]
    assert recs[1]["amount_paisa"] == 25000
    assert recs[1]["fee_paisa"] == 500
    assert "REF_ABC_12345678" in recs[1]["utr"]


def test_xlsx_upload_and_executive_pack_e2e(client, engine):
    import openpyxl
    with Session(engine) as session:
        rev = Reviewer(email="reviewer_xlsx@test.com", hashed_password=get_password_hash("pw"), role=Role.REVIEWER, portfolio_id="PORTFOLIO_A")
        case = ReconciliationCase(
            case_id="case_xlsx_e2e",
            portfolio_id="PORTFOLIO_A",
            exception_code="GATEWAY_FEE_MISMATCH",
            severity="HIGH",
            status=CaseStatus.OPEN,
            expected_paisa=1250000,
            actual_paisa=1220500,
            delta_paisa=29500,
            confidence_score=0.95,
            opened_at=utc_now(),
            explanation="Gateway fee difference detected",
            suggested_action="Request gateway settlement sheet"
        )
        session.add_all([rev, case])
        session.commit()

    # Login as reviewer
    res_login = client.post("/auth/login", data={"username": "reviewer_xlsx@test.com", "password": "pw"})
    cookie = res_login.cookies.get("session_token")

    # Generate XLSX spreadsheet
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dispute_Proof"
    ws.append(["UTR / Reference #", "Gross Amount", "MDR", "GST", "Net Payout", "Settlement Date"])
    ws.append(["AURA_DISPUTE_8811", 12500.00, 250.00, 45.00, 12205.00, "2026-09-02"])

    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    # Upload XLSX evidence
    res_upload = client.post(
        "/api/exceptions/case_xlsx_e2e/evidence/upload",
        files={"file": ("dispute_resolution.xlsx", io.BytesIO(xlsx_bytes), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers={"X-CSRF-Protection": "1"},
        cookies={"session_token": cookie}
    )
    assert res_upload.status_code == 200
    upload_data = res_upload.json()
    assert upload_data["status"] == "staged"
    assert upload_data["file_type"] == "XLSX"
    assert upload_data["summary"]["total_records"] == 1
    assert upload_data["summary"]["total_amount_paisa"] == 1250000
    assert upload_data["summary"]["net_amount_paisa"] == 1220500

    # Ensure records are NOT DOC_REF_EXTRACTED with 0 paisa!
    recs = upload_data["extracted_records"]
    assert len(recs) == 1
    assert recs[0]["utr"] == "AURA_DISPUTE_8811"
    assert recs[0]["amount_paisa"] == 1250000
    assert recs[0]["net_paisa"] == 1220500

    # Submit review to permanently commit proof
    res_review = client.post(
        "/api/exceptions/case_xlsx_e2e/review",
        json={"action": "APPROVE", "reason": "Verified bank credit advice spreadsheet proof."},
        headers={"X-CSRF-Protection": "1", "Idempotency-Key": "test-idem-xlsx-1"},
        cookies={"session_token": cookie}
    )
    assert res_review.status_code == 200

    # Verify Executive Pack has committed XLSX file and extracted rows
    res_exec = client.get("/api/exceptions/case_xlsx_e2e/executive-pack", cookies={"session_token": cookie})
    assert res_exec.status_code == 200
    pack = res_exec.json()
    attached_files = pack["evidence_pack"]["attached_evidence_files"]
    assert len(attached_files) == 1
    assert attached_files[0]["file_type"] == "XLSX"
    assert attached_files[0]["filename"] == "dispute_resolution.xlsx"
    extracted_txns = pack["evidence_pack"]["extracted_transactions"]
    assert len(extracted_txns) == 1
    assert extracted_txns[0]["utr"] == "AURA_DISPUTE_8811"
    assert extracted_txns[0]["amount_paisa"] == 1250000
    assert extracted_txns[0]["net_paisa"] == 1220500

def test_document_parser_unique_transaction_reference_and_utr():
    slip_text = """AURA BANK LIMITED
Date of Transfer: 02-Sep-2026
Advice Reference: ADV-AURA-2026-992211
Unique Transaction Reference (UTR): CMS998811223344
Sender: TECH CORP INDIA
Beneficiary Account: 50200012345678
Amount: INR 1,25,000.00
Charges: INR 250.00
GST (18%): INR 45.00
Net Credited Amount: INR 1,24,705.00
Remarks: Vendor monthly reconciliation payout"""
    res = parse_document_text(slip_text)
    assert len(res["records"]) == 1
    rec = res["records"][0]
    assert rec["utr"] == "CMS998811223344"
    assert rec["amount_paisa"] == 12500000
    assert rec["fee_paisa"] == 25000
    assert rec["tax_paisa"] == 4500
    assert rec["net_paisa"] == 12470500
    assert rec["timestamp"] == "2026-09-02"

def test_document_parser_end_to_end_id_and_payment_ref():
    slip_text = """APEX BANK TRANSACTION CONFIRMATION
End-to-End ID: E2E-APEX-44332211
Payment Reference Number: PAY-REF-88776655
Transaction Amount: 85,000.00
Value Date: 2026-09-01
Status: SUCCESS"""
    res = parse_document_text(slip_text)
    assert len(res["records"]) == 1
    rec = res["records"][0]
    assert rec["utr"] == "E2E-APEX-44332211"
    assert rec["amount_paisa"] == 8500000
    assert rec["timestamp"] == "2026-09-01"

def test_document_parser_two_line_layout():
    slip_text = """NOVA BANK DEBIT ADVICE
Bank Reference No:
NOVA_REF_776655
Total Paid: Rs. 15,400.00
Payment Date: 01/09/2026"""
    res = parse_document_text(slip_text)
    assert len(res["records"]) == 1
    rec = res["records"][0]
    assert rec["utr"] == "NOVA_REF_776655"
    assert rec["amount_paisa"] == 1540000
    assert rec["timestamp"] == "2026-09-01"

def test_document_parser_remittance_and_transaction_id():
    slip_text = """REMITTANCE ADVICE SLIP
Remittance Reference: REM/2026/09/554433
Amount Paid: $4,500.00
Date: September 02, 2026
Narration: Supplier invoice reimbursement"""
    res = parse_document_text(slip_text)
    assert len(res["records"]) == 1
    rec = res["records"][0]
    assert rec["utr"] == "REM/2026/09/554433"
    assert rec["amount_paisa"] == 450000
    assert rec["timestamp"] == "2026-09-02"

def test_document_parser_genuinely_no_reference_fallback():
    slip_text = """CASH PAYMENT CONFIRMATION RECEIPT
Total Amount: INR 500.00
Date: 2026-09-02
Office stationary supplies paid in cash without transaction IDs"""
    res = parse_document_text(slip_text)
    assert len(res["records"]) == 1
    rec = res["records"][0]
    # Uses DOC_REF_EXTRACTED only when document genuinely contains no reference
    assert rec["utr"] == "DOC_REF_EXTRACTED"
    assert rec["amount_paisa"] == 50000

def test_pdf_advice_slip_binary_upload_e2e(client, engine):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    with Session(engine) as session:
        rev = Reviewer(email="reviewer_pdf@test.com", hashed_password=get_password_hash("pw"), role=Role.REVIEWER, portfolio_id="PORTFOLIO_A")
        case = ReconciliationCase(
            case_id="case_pdf_e2e",
            portfolio_id="PORTFOLIO_A",
            exception_code="SETTLEMENT_DEFICIT",
            severity="HIGH",
            status=CaseStatus.OPEN,
            expected_paisa=3500000,
            actual_paisa=3400000,
            delta_paisa=100000,
            confidence_score=0.95,
            opened_at=utc_now(),
            explanation="Settlement deficit dispute",
            suggested_action="Verify payment advice slip PDF"
        )
        session.add_all([rev, case])
        session.commit()

    # Login
    res_login = client.post("/auth/login", data={"username": "reviewer_pdf@test.com", "password": "pw"})
    cookie = res_login.cookies.get("session_token")

    # Generate a real binary PDF advice slip
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.drawString(100, 720, "PAYMENT ADVICE SLIP")
    c.drawString(100, 700, "Transaction ID: TXN_PDF_CONFIRM_9988")
    c.drawString(100, 680, "Advice Reference: ADV-2026-PDF-01")
    c.drawString(100, 660, "Gross Amount: INR 35,000.00")
    c.drawString(100, 640, "Fee: INR 100.00")
    c.drawString(100, 620, "Net Amount: INR 34,900.00")
    c.drawString(100, 600, "Date: 02-Sep-2026")
    c.drawString(100, 580, "Beneficiary Account: 50100987654321")
    c.drawString(100, 560, "Remarks: Full settlement confirmation")
    c.save()
    pdf_bytes = buf.getvalue()

    # Upload PDF advice slip
    res_upload = client.post(
        "/api/exceptions/case_pdf_e2e/evidence/upload",
        files={"file": ("payment_advice_slip.pdf", io.BytesIO(pdf_bytes), "application/pdf")},
        headers={"X-CSRF-Protection": "1"},
        cookies={"session_token": cookie}
    )
    assert res_upload.status_code == 200
    upload_data = res_upload.json()
    assert upload_data["status"] == "staged"
    assert upload_data["file_type"] == "PDF"
    assert upload_data["summary"]["total_records"] == 1
    assert upload_data["summary"]["total_amount_paisa"] == 3500000
    assert upload_data["summary"]["net_amount_paisa"] == 3490000

    recs = upload_data["extracted_records"]
    assert len(recs) == 1
    assert recs[0]["utr"] == "TXN_PDF_CONFIRM_9988"
    assert recs[0]["amount_paisa"] == 3500000
    assert recs[0]["fee_paisa"] == 10000
    assert recs[0]["net_paisa"] == 3490000
    assert recs[0]["timestamp"] == "2026-09-02"

    # Submit review to permanently commit proof
    res_review = client.post(
        "/api/exceptions/case_pdf_e2e/review",
        json={"action": "APPROVE", "reason": "Verified bank advice slip PDF with Transaction ID."},
        headers={"X-CSRF-Protection": "1", "Idempotency-Key": "test-idem-pdf-1"},
        cookies={"session_token": cookie}
    )
    assert res_review.status_code == 200

    # Verify Executive Pack has committed PDF file and extracted records
    res_exec = client.get("/api/exceptions/case_pdf_e2e/executive-pack", cookies={"session_token": cookie})
    assert res_exec.status_code == 200
    pack = res_exec.json()
    attached_files = pack["evidence_pack"]["attached_evidence_files"]
    assert len(attached_files) == 1
    assert attached_files[0]["file_type"] == "PDF"
    assert attached_files[0]["filename"] == "payment_advice_slip.pdf"
    extracted_txns = pack["evidence_pack"]["extracted_transactions"]
    assert len(extracted_txns) == 1
    assert extracted_txns[0]["utr"] == "TXN_PDF_CONFIRM_9988"
    assert extracted_txns[0]["amount_paisa"] == 3500000
    assert extracted_txns[0]["net_paisa"] == 3490000

def test_bai2_multiline_continuation_88():
    bai2_txt = """01,CITIUS33,LEDGERLENS,260902,1430,001,80,1,2/
02,LEDGERLENS,CITIUS33,1,260902,1430,INR,2/
03,50100987654321,INR,010,5000000,,,/
16,175,25000,V,CITI_UTR_9988,CUST_REF_01,Dispute settlement payment
88,credit for merchant account M12345/
16,475,500,0,NONREF,E2E_DEBIT_5544,Chargeback processing fee/
49,24500,2/
98,24500,1,2/
99,24500,1,2/"""
    res = parse_bai2(bai2_txt)
    assert len(res["records"]) == 2
    rec1 = res["records"][0]
    assert rec1["utr"] == "CITI_UTR_9988"
    assert rec1["amount_paisa"] == 2500000
    assert "Dispute settlement payment credit for merchant account M12345" in rec1["narrative"]
    assert rec1["timestamp"] == "2026-09-02"

    rec2 = res["records"][1]
    assert rec2["utr"] == "E2E_DEBIT_5544"
    assert rec2["amount_paisa"] == 50000
    assert "[DBIT]" in rec2["narrative"]

def test_bai2_single_line_streaming_slashes():
    bai2_stream = "01,BOFAUS3N,LEDGERLENS,260902,1200/02,LEDGERLENS,BOFAUS3N,1,260902/03,1122334455,USD/16,175,10000,0,BOFA_REF_7766,BOFA_CUST_1,Merchant wire credit/49,10000,1/98,10000,1,1/99,10000,1,1/"
    res = parse_bai2(bai2_stream)
    assert len(res["records"]) == 1
    rec = res["records"][0]
    assert rec["utr"] == "BOFA_REF_7766"
    assert rec["amount_paisa"] == 1000000
    assert rec["timestamp"] == "2026-09-02"

def test_bai2_customer_reference_fallback():
    bai2_txt = """01,HSBCINBB,LEDGERLENS,260902,0900,1/
02,LEDGERLENS,HSBCINBB,1,260902/
03,9988776655,INR/
16,175,15000,0,NONREF,HSBC_E2E_443322,ACH Vendor settlement/
49,15000,1/
98,15000,1,1/
99,15000,1,1/"""
    res = parse_bai2(bai2_txt)
    assert len(res["records"]) == 1
    rec = res["records"][0]
    assert rec["utr"] == "HSBC_E2E_443322"
    assert rec["amount_paisa"] == 1500000

def test_bai2_malformed_file_validation_error():
    import pytest
    with pytest.raises(ValueError, match="Malformed BAI2 file"):
        parse_bai2("CORRUPTED TEXT FILE WITHOUT ANY 01 OR 16 BAI2 RECORDS")

    with pytest.raises(ValueError, match="Malformed BAI2 file"):
        parse_bai2("   \n   \t  ")

def test_bai2_e2e_upload_and_executive_pack(client, engine):
    with Session(engine) as session:
        rev = Reviewer(email="reviewer_bai2@test.com", hashed_password=get_password_hash("pw"), role=Role.REVIEWER, portfolio_id="PORTFOLIO_A")
        case = ReconciliationCase(
            case_id="case_bai2_e2e",
            portfolio_id="PORTFOLIO_A",
            exception_code="SETTLEMENT_DEFICIT",
            severity="HIGH",
            status=CaseStatus.OPEN,
            expected_paisa=5000000,
            actual_paisa=4900000,
            delta_paisa=100000,
            confidence_score=0.96,
            opened_at=utc_now(),
            explanation="Settlement deficit dispute",
            suggested_action="Verify BAI2 cash management statement"
        )
        session.add_all([rev, case])
        session.commit()

    # Login
    res_login = client.post("/auth/login", data={"username": "reviewer_bai2@test.com", "password": "pw"})
    cookie = res_login.cookies.get("session_token")

    bai2_content = """01,JPMCUS33,LEDGERLENS,260902,1500,1/
02,LEDGERLENS,JPMCUS33,1,260902/
03,4455667788,INR/
16,175,50000,0,JPMC_LIVE_BAI_9900,JPMC_CUST_REF,JPMorgan dispute credit settlement/
49,50000,1/
98,50000,1,1/
99,50000,1,1/""".encode("utf-8")

    # 1. Test malformed .bai2 upload returns clear 400 validation error
    res_bad = client.post(
        "/api/exceptions/case_bai2_e2e/evidence/upload",
        files={"file": ("corrupted_feed.bai2", io.BytesIO(b"MALFORMED_GARBAGE_NO_BAI2_RECORDS"), "text/plain")},
        headers={"X-CSRF-Protection": "1"},
        cookies={"session_token": cookie}
    )
    assert res_bad.status_code == 400
    assert "Malformed BAI2 file" in res_bad.json()["detail"]

    # 2. Upload valid .bai2 file
    res_upload = client.post(
        "/api/exceptions/case_bai2_e2e/evidence/upload",
        files={"file": ("commercial_bank_stmt.bai2", io.BytesIO(bai2_content), "text/plain")},
        headers={"X-CSRF-Protection": "1"},
        cookies={"session_token": cookie}
    )
    assert res_upload.status_code == 200
    upload_data = res_upload.json()
    assert upload_data["status"] == "staged"
    assert upload_data["file_type"] == "BAI2"
    assert upload_data["summary"]["total_records"] == 1
    assert upload_data["summary"]["total_amount_paisa"] == 5000000
    assert upload_data["summary"]["net_amount_paisa"] == 5000000

    recs = upload_data["extracted_records"]
    assert len(recs) == 1
    assert recs[0]["utr"] == "JPMC_LIVE_BAI_9900"
    assert recs[0]["amount_paisa"] == 5000000
    assert recs[0]["timestamp"] == "2026-09-02"

    # 3. Submit review to permanently commit proof
    res_review = client.post(
        "/api/exceptions/case_bai2_e2e/review",
        json={"action": "APPROVE", "reason": "Verified bank statement in BAI2 format."},
        headers={"X-CSRF-Protection": "1", "Idempotency-Key": "test-idem-bai2-1"},
        cookies={"session_token": cookie}
    )
    assert res_review.status_code == 200

    # 4. Verify Executive Pack has committed BAI2 file and extracted records
    res_exec = client.get("/api/exceptions/case_bai2_e2e/executive-pack", cookies={"session_token": cookie})
    assert res_exec.status_code == 200
    pack = res_exec.json()
    attached_files = pack["evidence_pack"]["attached_evidence_files"]
    assert len(attached_files) == 1
    assert attached_files[0]["file_type"] == "BAI2"
    assert attached_files[0]["filename"] == "commercial_bank_stmt.bai2"
    extracted_txns = pack["evidence_pack"]["extracted_transactions"]
    assert len(extracted_txns) == 1
    assert extracted_txns[0]["utr"] == "JPMC_LIVE_BAI_9900"
    assert extracted_txns[0]["amount_paisa"] == 5000000

def test_bai2_account_status_03_records():
    bai2_content = """01,123456789,20260904,1000,2,
02,IN00BANK000000000001,123456789,INR,1000,
03,IN00BANK000000000001,010,1500,1,TXN-2026-00201,Alpha Electronics,
03,IN00BANK000000000001,010,825,1,TXN-2026-00202,Metro Travel,
03,IN00BANK000000000001,010,450,1,TXN-2026-00203,Cloud Services Ltd,
03,IN00BANK000000000001,010,2750,1,TXN-2026-00204,Office Mart,
49,5525,4,
98,5525,1,8,
99,5525,1,8,"""
    res = parse_bai2(bai2_content)
    assert res["summary"]["total_records"] == 4
    assert res["summary"]["total_amount_paisa"] == 552500
    recs = res["records"]
    assert recs[0]["utr"] == "TXN-2026-00201"
    assert recs[0]["amount_paisa"] == 150000
    assert recs[0]["timestamp"] == "2026-09-04"
    assert "Alpha Electronics" in recs[0]["narrative"]

    assert recs[1]["utr"] == "TXN-2026-00202"
    assert recs[1]["amount_paisa"] == 82500
    assert "Metro Travel" in recs[1]["narrative"]

    assert recs[2]["utr"] == "TXN-2026-00203"
    assert recs[2]["amount_paisa"] == 45000

    assert recs[3]["utr"] == "TXN-2026-00204"
    assert recs[3]["amount_paisa"] == 275000


def test_evidence_upload_luhn_pan_sanitization_in_storage(client, engine):
    """
    Verify that when evidence containing a valid Luhn PAN is uploaded via the real ingestion endpoint
    (/api/exceptions/{case_id}/evidence/upload), the PAN is masked before being persisted in
    EvidenceAttachment.extracted_data_json.
    """
    with Session(engine) as session:
        rev = Reviewer(
            email="reviewer_pan_test@test.com",
            hashed_password=get_password_hash("pw"),
            role=Role.REVIEWER,
            portfolio_id="PORTFOLIO_PAN"
        )
        case = ReconciliationCase(
            case_id="case_pan_sanitization",
            portfolio_id="PORTFOLIO_PAN",
            exception_code="FEE_RATE_MISMATCH",
            severity="MEDIUM",
            status=CaseStatus.OPEN,
            expected_paisa=100000,
            actual_paisa=98000,
            delta_paisa=2000,
            confidence_score=0.95,
            opened_at=utc_now(),
            explanation="PAN sanitization test case",
            suggested_action="Verify PAN masking in storage"
        )
        session.add_all([rev, case])
        session.commit()

    # Login
    res_login = client.post("/auth/login", data={"username": "reviewer_pan_test@test.com", "password": "pw"})
    cookie = res_login.cookies.get("session_token")

    # Raw CSV containing valid Luhn Visa PAN: 4532015012345671
    valid_luhn_pan = "4532015012345671"
    raw_csv = (
        "UTR,Amount,Fee,Tax,Net_Amount,Date,Customer_Card,Notes\n"
        f"UTR_PAN_001,1000.00,20.00,3.60,976.40,2026-09-01,{valid_luhn_pan},Customer card payment with {valid_luhn_pan}\n"
    ).encode("utf-8")

    res_upload = client.post(
        "/api/exceptions/case_pan_sanitization/evidence/upload",
        files={"file": ("customer_transactions.csv", io.BytesIO(raw_csv), "text/csv")},
        headers={"X-CSRF-Protection": "1"},
        cookies={"session_token": cookie}
    )
    assert res_upload.status_code == 200
    upload_resp = res_upload.json()
    assert upload_resp["status"] == "staged"

    # Verify that in the API response, raw PAN is masked
    response_text = json.dumps(upload_resp)
    assert valid_luhn_pan not in response_text
    assert "XXXX-XXXX-XXXX-5671" in response_text

    # Verify that in the database storage (EvidenceAttachment.extracted_data_json), raw PAN was NEVER stored
    with Session(engine) as session:
        att = session.exec(
            select(EvidenceAttachment).where(EvidenceAttachment.case_id == "case_pan_sanitization")
        ).first()
        assert att is not None
        assert att.extracted_data_json is not None
        # Assert raw PAN is NOT in the database JSON
        assert valid_luhn_pan not in att.extracted_data_json
        # Assert masked format IS in the database JSON
        assert "XXXX-XXXX-XXXX-5671" in att.extracted_data_json





